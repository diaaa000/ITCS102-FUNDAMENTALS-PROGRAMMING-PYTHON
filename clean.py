import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

# Create Excel File if not exists
if not os.path.exists("SchoolEvent_Registration.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Registrations"
    ws.append(["ID", "Student Name", "Course", "Year Level", "Event Name"])
    wb.save("SchoolEvent_Registration.xlsx")

# GUI Setup
window = tk.Tk()
window.title("Event Registration")
window.geometry("900x600")
window.configure(bg="pale violet red")

# Variables
id_var = tk.IntVar()
name_var = tk.StringVar()
course_var = tk.StringVar()
year_var = tk.StringVar()
event_var = tk.StringVar()

# Frame
frame = tk.LabelFrame(
    window,
    text="Event Registration Form",
    bg="maroon",
    fg="white",
    padx=10,
    pady=10
)
frame.pack(fill="x", padx=10, pady=10)

# Labels & Entries
tk.Label(frame, text="Registration ID", font=("Poppins", 12), bg="maroon", fg="white").grid(row=0, column=0, padx=5, pady=5)
Rid_entry = tk.Entry(frame, font=("Poppins", 12), textvariable=id_var, state="readonly")
Rid_entry.grid(row=0, column=1)

tk.Label(frame, text="Student Name", font=("Poppins", 12), bg="maroon", fg="white").grid(row=1, column=0, padx=5, pady=5)
Sname_entry = tk.Entry(frame, font=("Poppins", 12), textvariable=name_var, width=30)
Sname_entry.grid(row=1, column=1)

tk.Label(frame, text="Course", font=("Poppins", 12), bg="maroon", fg="white").grid(row=2, column=0, padx=5, pady=5)
Course_entry = tk.Entry(frame, textvariable=course_var, width=30)
Course_entry.grid(row=2, column=1)

tk.Label(frame, text="Year Level", font=("Poppins", 12), bg="maroon", fg="white").grid(row=3, column=0, padx=5, pady=5)
year_combo = ttk.Combobox(frame, textvariable=year_var, values=["1st Year", "2nd Year", "3rd Year", "4th Year"])
year_combo.grid(row=3, column=1)

tk.Label(frame, text="Event Name", font=("Poppins", 12), bg="maroon", fg="white").grid(row=4, column=0, padx=5, pady=5)
Ename_entry = tk.Entry(frame, textvariable=event_var, width=30)
Ename_entry.grid(row=4, column=1)

# Functions
def generate_id():
    wb = load_workbook("SchoolEvent_Registration.xlsx")
    ws = wb["Registrations"]
    if ws.max_row == 1:
        return 1
    last_id = ws.cell(ws.max_row, 1).value
    return last_id + 1

def clear_all():
    id_var.set(generate_id())
    name_var.set("")
    course_var.set("")
    year_var.set("")
    event_var.set("")

def load_data():
    for row in tree.get_children():
        tree.delete(row)
    wb = load_workbook("SchoolEvent_Registration.xlsx")
    ws = wb["Registrations"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", tk.END, values=row)

def add_record():
    if not name_var.get() or not course_var.get() or not year_var.get() or not event_var.get():
        messagebox.showerror("Oops! Wrong Input", "Please fill in all fields.")
        return
    wb = load_workbook("SchoolEvent_Registration.xlsx")
    ws = wb["Registrations"]
    ws.append([id_var.get(), name_var.get(), course_var.get(), year_var.get(), event_var.get()])
    wb.save("SchoolEvent_Registration.xlsx")
    messagebox.showinfo("Yay! Success", "Your registration added successfully!")
    load_data()
    clear_all()

def update_record():
    selected = tree.selection()
    if not selected:
        messagebox.showerror("Error", "Select a record first.")
        return
    confirm = messagebox.askyesno("Confirm Update", "Do you want to update this record?")
    if not confirm:
        return
    wb = load_workbook("SchoolEvent_Registration.xlsx")
    ws = wb["Registrations"]
    selected_id = int(id_var.get())
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == selected_id:
            ws.cell(row, 2).value = name_var.get()
            ws.cell(row, 3).value = course_var.get()
            ws.cell(row, 4).value = year_var.get()
            ws.cell(row, 5).value = event_var.get()
            break
    wb.save("SchoolEvent_Registration.xlsx")
    messagebox.showinfo("Yay! Success", "Record updated successfully!")
    load_data()
    clear_all()

def delete_record():
    selected = tree.selection()
    if not selected:
        messagebox.showerror("Oops! Error", "Please select a record first.")
        return
    confirm = messagebox.askyesno("Confirm Delete", "Do you want to delete this record?")
    if not confirm:
        return
    wb = load_workbook("SchoolEvent_Registration.xlsx")
    ws = wb["Registrations"]
    selected_id = int(id_var.get())
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == selected_id:
            ws.delete_rows(row)
            break
    wb.save("SchoolEvent_Registration.xlsx")
    messagebox.showinfo("Deleted", "Record deleted successfully!")
    load_data()
    clear_all()

def select_record(event):
    selected = tree.focus()
    if not selected:
        return
    values = tree.item(selected, "values")
    id_var.set(values[0])
    name_var.set(values[1])
    course_var.set(values[2])
    year_var.set(values[3])
    event_var.set(values[4])

# Buttons
btn_frame = tk.Frame(window)
btn_frame.pack(pady=10)

tk.Button(btn_frame, text="Add", font=("Poppins",12,"bold"), bg="deep pink", width=12, command=add_record).grid(row=0, column=0, padx=5)
tk.Button(btn_frame, text="Update", font=("Poppins",12,"bold"), bg="deep pink", width=12, command=update_record).grid(row=0, column=1, padx=5)
tk.Button(btn_frame, text="Delete", font=("Poppins",12,"bold"), bg="deep pink", width=12, command=delete_record).grid(row=0, column=2, padx=5)
tk.Button(btn_frame, text="Clear", font=("Poppins",12,"bold"), bg="deep pink", width=12, command=clear_all).grid(row=0, column=3, padx=5)

# TreeView
table_frame = tk.Frame(window)
table_frame.pack(fill="both", expand=True, padx=10, pady=10)

columns = ("ID", "Student Name", "Course", "Year Level", "Event Name")
tree = ttk.Treeview(table_frame, columns=columns, show="headings")
for col in columns:
    tree.heading(col, text=col)
    tree.column(col, width=150)
tree.pack(fill="both", expand=True)
tree.bind("<<TreeviewSelect>>", select_record)

# Startup
id_var.set(generate_id())
load_data()

window.mainloop()
