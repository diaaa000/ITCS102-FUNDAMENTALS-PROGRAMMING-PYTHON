import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

# ==========================
# CREATE EXCEL FILE
# ==========================

FILE_NAME = "event_registration.xlsx"

if not os.path.exists(FILE_NAME):
    wb = Workbook()
    ws = wb.active
    ws.title = "Registrations"
    ws.append([
        "ID",
        "Student Name",
        "Course",
        "Year Level",
        "Event Name"
    ])
    wb.save(FILE_NAME)

# ==========================
# FUNCTIONS
# ==========================

def generate_id():
    wb = load_workbook(FILE_NAME)
    ws = wb["Registrations"]

    if ws.max_row == 1:
        return 1

    last_id = ws.cell(ws.max_row, 1).value
    return last_id + 1


def clear_fields():
    id_var.set(generate_id())
    name_var.set("")
    course_var.set("")
    year_var.set("")
    event_var.set("")


def load_data():
    for row in tree.get_children():
        tree.delete(row)

    wb = load_workbook(FILE_NAME)
    ws = wb["Registrations"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", tk.END, values=row)


def add_record():

    if name_var.get() == "" or course_var.get() == "" or \
       year_var.get() == "" or event_var.get() == "":
        messagebox.showerror(
            "Input Error",
            "Please fill in all fields."
        )
        return

    wb = load_workbook(FILE_NAME)
    ws = wb["Registrations"]

    ws.append([
        id_var.get(),
        name_var.get(),
        course_var.get(),
        year_var.get(),
        event_var.get()
    ])

    wb.save(FILE_NAME)

    messagebox.showinfo(
        "Success",
        "Registration added successfully!"
    )

    load_data()
    clear_fields()


def update_record():

    selected = tree.selection()

    if not selected:
        messagebox.showerror(
            "Error",
            "Select a record first."
        )
        return

    confirm = messagebox.askyesno(
        "Confirm Update",
        "Update this record?"
    )

    if not confirm:
        return

    wb = load_workbook(FILE_NAME)
    ws = wb["Registrations"]

    selected_id = int(id_var.get())

    for row in range(2, ws.max_row + 1):

        if ws.cell(row, 1).value == selected_id:

            ws.cell(row, 2).value = name_var.get()
            ws.cell(row, 3).value = course_var.get()
            ws.cell(row, 4).value = year_var.get()
            ws.cell(row, 5).value = event_var.get()

            break

    wb.save(FILE_NAME)

    messagebox.showinfo(
        "Success",
        "Record updated successfully!"
    )

    load_data()
    clear_fields()


def delete_record():

    selected = tree.selection()

    if not selected:
        messagebox.showerror(
            "Error",
            "Select a record first."
        )
        return

    confirm = messagebox.askyesno(
        "Confirm Delete",
        "Delete this record?"
    )

    if not confirm:
        return

    wb = load_workbook(FILE_NAME)
    ws = wb["Registrations"]

    selected_id = int(id_var.get())

    for row in range(2, ws.max_row + 1):

        if ws.cell(row, 1).value == selected_id:
            ws.delete_rows(row)
            break

    wb.save(FILE_NAME)

    messagebox.showinfo(
        "Deleted",
        "Record deleted successfully!"
    )

    load_data()
    clear_fields()


def select_record(event):

    selected = tree.focus()

    if not selected:
        return

    values = tree.item(selected, "values")

    id_var.set(values[0])
    name_var.set(values[1])
    course_var.set(values[2])
    year_var.set(values[3])
    event_var.set(values[4])


# ==========================
# GUI
# ==========================

root = tk.Tk()
root.title("School Event Registration System")
root.geometry("900x600")

# Variables
id_var = tk.IntVar()
name_var = tk.StringVar()
course_var = tk.StringVar()
year_var = tk.StringVar()
event_var = tk.StringVar()

# ==========================
# FORM FRAME
# ==========================

form_frame = tk.LabelFrame(
    root,
    text="Event Registration Form",
    padx=10,
    pady=10
)

form_frame.pack(fill="x", padx=10, pady=10)

tk.Label(form_frame, text="Registration ID").grid(
    row=0, column=0, padx=5, pady=5
)

tk.Entry(
    form_frame,
    textvariable=id_var,
    state="readonly"
).grid(row=0, column=1)

tk.Label(form_frame, text="Student Name").grid(
    row=1, column=0, padx=5, pady=5
)

tk.Entry(
    form_frame,
    textvariable=name_var,
    width=30
).grid(row=1, column=1)

tk.Label(form_frame, text="Course").grid(
    row=2, column=0, padx=5, pady=5
)

tk.Entry(
    form_frame,
    textvariable=course_var,
    width=30
).grid(row=2, column=1)

tk.Label(form_frame, text="Year Level").grid(
    row=3, column=0, padx=5, pady=5
)

year_combo = ttk.Combobox(
    form_frame,
    textvariable=year_var,
    values=[
        "1st Year",
        "2nd Year",
        "3rd Year",
        "4th Year"
    ]
)

year_combo.grid(row=3, column=1)

tk.Label(form_frame, text="Event Name").grid(
    row=4, column=0, padx=5, pady=5
)

tk.Entry(
    form_frame,
    textvariable=event_var,
    width=30
).grid(row=4, column=1)

# ==========================
# BUTTONS
# ==========================

btn_frame = tk.Frame(root)
btn_frame.pack(pady=10)

tk.Button(
    btn_frame,
    text="Add",
    width=12,
    command=add_record
).grid(row=0, column=0, padx=5)

tk.Button(
    btn_frame,
    text="Update",
    width=12,
    command=update_record
).grid(row=0, column=1, padx=5)

tk.Button(
    btn_frame,
    text="Delete",
    width=12,
    command=delete_record
).grid(row=0, column=2, padx=5)

tk.Button(
    btn_frame,
    text="Clear",
    width=12,
    command=clear_fields
).grid(row=0, column=3, padx=5)

# ==========================
# TREEVIEW
# ==========================

table_frame = tk.Frame(root)
table_frame.pack(fill="both", expand=True, padx=10, pady=10)

columns = (
    "ID",
    "Student Name",
    "Course",
    "Year Level",
    "Event Name"
)

tree = ttk.Treeview(
    table_frame,
    columns=columns,
    show="headings"
)

for col in columns:
    tree.heading(col, text=col)
    tree.column(col, width=150)

tree.pack(fill="both", expand=True)

tree.bind("<<TreeviewSelect>>", select_record)

# ==========================
# STARTUP
# ==========================

id_var.set(generate_id())
load_data()

root.mainloop()